"""Station-scoped reporting. Exports use the same data as the on-screen reports."""
import csv, io
from datetime import date, timedelta
from decimal import Decimal
from typing import Literal
from html import escape
from fastapi import APIRouter, HTTPException, Query, Response
from sqlalchemy import select, func
from .main import DB, Actor, Shift, Tank, Product, Nozzle, Expense, Delivery, User, scoped, permit
from .operations import StockMovement, TankReading, Customer, CustomerPayment, Maintenance, Equipment, Incident, BusinessDay, customer_totals, today, ZERO
router=APIRouter()
ReportKind=Literal['sales','inventory','expenses','deliveries','credit','maintenance','daily','management']

def day_summary(db,station_id,day):
    rows=db.scalars(select(Shift).where(Shift.station_id==station_id,Shift.business_date==day,Shift.status=='reconciled')).all()
    expenses=db.scalars(select(Expense).where(Expense.station_id==station_id,Expense.business_date==day,Expense.status=='approved')).all()
    receipts=db.scalars(select(CustomerPayment).where(CustomerPayment.station_id==station_id,CustomerPayment.business_date==day)).all()
    total=lambda k:sum((getattr(s,k) for s in rows),ZERO)
    return {'date':day,'litres':total('litres'),'expected':total('expected'),'cash':total('cash'),'momo':total('momo'),'card':total('card'),'credit_sales':total('credit'),'fuel_collections':total('cash')+total('momo')+total('card'),'accounted_sales':total('cash')+total('momo')+total('card')+total('credit'),'variance':total('variance'),'customer_debt_payments':sum((p.amount for p in receipts),ZERO),'expenses':sum((e.amount for e in expenses),ZERO)}

def inventory_rows(db,station_id,start,end):
    rows=[]
    for tank in db.scalars(select(Tank).where(Tank.station_id==station_id)).all():
        movements=db.scalars(select(StockMovement).where(StockMovement.tank_id==tank.id,StockMovement.business_date<=end)).all()
        before=sum((m.quantity for m in movements if m.business_date<start),ZERO)
        period=[m for m in movements if start<=m.business_date<=end]
        opening=before+sum((m.quantity for m in period if m.source in ['opening','migration_opening']),ZERO)
        deliveries=sum((m.quantity for m in period if m.source=='delivery'),ZERO)
        sales=-sum((m.quantity for m in period if m.source=='shift'),ZERO)
        adjustments=sum((m.quantity for m in period if m.source=='adjustment'),ZERO)
        closing=opening+deliveries-sales+adjustments
        reading=db.scalar(select(TankReading).where(TankReading.tank_id==tank.id,TankReading.business_date==end).order_by(TankReading.id.desc()))
        rows.append({'tank_id':tank.id,'tank':tank.name,'fuel':db.get(Product,tank.product_id).name,'opening':opening,'deliveries':deliveries,'litres_sold':sales,'adjustments':adjustments,'expected_closing':closing,'physical_closing':reading.physical if reading else None,'variance':reading.physical-closing if reading else None})
    return rows

def report_data(db,station_id,kind,start,end,nozzle_id=None,attendant_id=None,product_id=None,payment_method=None):
    if start>end or (end-start).days>366:raise HTTPException(422,'Choose a valid date range of no more than 366 days')
    if kind=='sales':
        stmt=select(Shift,Nozzle,Product,User).join(Nozzle,Nozzle.id==Shift.nozzle_id).join(Tank,Tank.id==Nozzle.tank_id).join(Product,Product.id==Tank.product_id).join(User,User.id==Shift.attendant_id).where(Shift.station_id==station_id,Shift.business_date>=start,Shift.business_date<=end,Shift.status=='reconciled')
        if nozzle_id:stmt=stmt.where(Shift.nozzle_id==nozzle_id)
        if attendant_id:stmt=stmt.where(Shift.attendant_id==attendant_id)
        if product_id:stmt=stmt.where(Product.id==product_id)
        if payment_method:stmt=stmt.where(getattr(Shift,payment_method)>0)
        rows=[{'date':s.business_date,'shift':s.id,'nozzle':n.name,'attendant':u.name,'fuel':p.name,'litres':s.litres,'price':s.price,'expected':s.expected,'cash':s.cash,'momo':s.momo,'card':s.card,'credit':s.credit,'variance':s.variance} for s,n,p,u in db.execute(stmt.order_by(Shift.business_date,Shift.id)).all()]
    elif kind=='inventory':rows=inventory_rows(db,station_id,start,end)
    elif kind=='expenses':rows=[{'date':e.business_date,'id':e.id,'description':e.description,'amount':e.amount,'status':e.status,'recorded_by':e.recorded_by,'approved_by':e.approved_by} for e in db.scalars(select(Expense).where(Expense.station_id==station_id,Expense.business_date>=start,Expense.business_date<=end).order_by(Expense.business_date,Expense.id)).all()]
    elif kind=='deliveries':rows=[{'date':d.business_date,'reference':d.reference,'supplier':d.supplier,'tank_id':d.tank_id,'waybill_litres':d.waybill_litres,'received_litres':d.received_litres,'variance_litres':d.received_litres-d.waybill_litres,'cost_per_litre':d.cost_per_litre,'total_cost':d.received_litres*d.cost_per_litre} for d in db.scalars(select(Delivery).where(Delivery.station_id==station_id,Delivery.business_date>=start,Delivery.business_date<=end).order_by(Delivery.business_date,Delivery.id)).all()]
    elif kind=='management':
        expenses=sum((e.amount for e in db.scalars(select(Expense).where(Expense.station_id==station_id,Expense.business_date>=start,Expense.business_date<=end,Expense.status=='approved')).all()),ZERO)
        sales=db.execute(select(Product.id,Product.name,func.sum(Shift.litres),func.sum(Shift.expected)).join(Tank,Tank.product_id==Product.id).join(Nozzle,Nozzle.tank_id==Tank.id).join(Shift,Shift.nozzle_id==Nozzle.id).where(Shift.station_id==station_id,Shift.business_date>=start,Shift.business_date<=end,Shift.status=='reconciled').group_by(Product.id,Product.name)).all()
        total_revenue=sum((r[3] for r in sales),ZERO);rows=[]
        for pid,name,litres,revenue in sales:
            delivery=db.scalar(select(Delivery).join(Tank,Tank.id==Delivery.tank_id).where(Tank.product_id==pid,Delivery.business_date<=end).order_by(Delivery.business_date.desc(),Delivery.id.desc()))
            price=delivery.cost_per_litre if delivery else None
            cost=(litres*price).quantize(Decimal('0.01')) if price is not None else None
            allocation=(expenses*revenue/total_revenue).quantize(Decimal('0.01')) if total_revenue else ZERO
            rows.append({'fuel':name,'litres':litres,'revenue':revenue,'reference_cost_per_litre':price,'estimated_fuel_cost':cost,'allocated_approved_expenses':allocation,'estimated_operating_margin':revenue-cost-allocation if cost is not None else None,'cost_basis':'Latest recorded delivery price; estimate' if delivery else 'Cost unavailable'})
    elif kind=='credit':
        rows=[{'customer':c.name,'credit_limit':c.credit_limit,**customer_totals(db,c.id)} for c in db.scalars(select(Customer).where(Customer.station_id==station_id)).all()]
    elif kind=='maintenance':rows=[{'id':m.id,'equipment':db.get(Equipment,m.equipment_id).name,'reported':m.created.date(),'fault':m.fault,'technician':m.technician,'status':m.status,'cost':m.cost,'next_service':m.next_service_date} for m in db.scalars(select(Maintenance).where(Maintenance.station_id==station_id,func.date(Maintenance.created)>=start,func.date(Maintenance.created)<=end)).all()]
    else:
        rows=[];day=start
        while day<=end:rows.append(day_summary(db,station_id,day));day+=timedelta(days=1)
    if len(rows)>20000:raise HTTPException(422,'Report too large; reduce the date range')
    numeric=[k for k in (rows[0] if rows else []) if any(isinstance(r.get(k),Decimal) for r in rows)]
    totals={k:sum((r.get(k) or ZERO for r in rows),ZERO) for k in numeric}
    # Price and point-in-time balances are not meaningful sums.
    for k in ['price','cost_per_litre','reference_cost_per_litre','credit_limit','opening','expected_closing','physical_closing','variance']:
        if kind=='inventory' or k in ['price','cost_per_litre','reference_cost_per_litre','credit_limit']:totals.pop(k,None)
    if kind=='management' and any(r['estimated_operating_margin'] is None for r in rows):
        totals.pop('estimated_operating_margin',None);totals.pop('estimated_fuel_cost',None)
    return {'kind':kind,'station_id':station_id,'start':start,'end':end,'basis':'Current customer account balances' if kind=='credit' else 'Estimated margin uses latest delivery prices and approved expenses; incomplete costs are left blank. This is not inventory valuation.' if kind=='management' else 'Business-date records; sales include reconciled shifts only','rows':rows,'totals':totals}

@router.get('/reports/{kind}')
def report(kind:ReportKind,station_id:int,db:DB,actor:Actor,start:date=Query(default_factory=today),end:date=Query(default_factory=today),nozzle_id:int | None=None,attendant_id:int | None=None,product_id:int | None=None,payment_method:Literal['cash','momo','card','credit'] | None=None):
    permit(actor,'owner','manager','supervisor','accountant','auditor');scoped(actor,station_id)
    return report_data(db,station_id,kind,start,end,nozzle_id,attendant_id,product_id,payment_method)

def safe_cell(v):
    if v is None:return ''
    if isinstance(v,(date,Decimal,int,float)):return v
    s=str(v)
    return "'"+s if s.lstrip().startswith(('=','+','-','@','\t','\r')) else s

def export_bytes(data,format):
    rows=data['rows'];headers=list(rows[0]) if rows else ['No records']
    if format=='csv':
        buffer=io.StringIO();writer=csv.writer(buffer);writer.writerow(headers)
        for row in rows:writer.writerow([safe_cell(row.get(h)) for h in headers])
        return buffer.getvalue().encode('utf-8-sig'),'text/csv; charset=utf-8'
    if format=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        workbook=Workbook();sheet=workbook.active;sheet.title='FuelFlow report';sheet.append(headers)
        for row in rows:sheet.append([safe_cell(row.get(h)) for h in headers])
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='007F76')
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for i,h in enumerate(headers,1):sheet.column_dimensions[get_column_letter(i)].width=min(45,max(16,len(h)+3))
        info=workbook.create_sheet('Report details');info.append(['FuelFlow Ghana',data['kind']]);info.append(['Start',data['start']]);info.append(['End',data['end']]);info.append(['Basis',data['basis']])
        for k,v in data['totals'].items():info.append([k,v])
        buffer=io.BytesIO();workbook.save(buffer);return buffer.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    buffer=io.BytesIO();styles=getSampleStyleSheet();style=styles['BodyText'];style.fontSize=6;style.leading=8
    doc=SimpleDocTemplate(buffer,pagesize=landscape(A4),leftMargin=24,rightMargin=24,topMargin=28,bottomMargin=28)
    story=[Paragraph('FuelFlow Ghana — '+escape(data['kind'].title())+' Report',styles['Heading1']),Paragraph(escape(f"{data['start']} to {data['end']} · {data['basis']}"),styles['Normal']),Spacer(1,15)]
    cells=[[Paragraph(escape(str(h).replace('_',' ')),style) for h in headers]]
    for row in rows:cells.append([Paragraph(escape(str(row.get(h) if row.get(h) is not None else '')),style) for h in headers])
    table=LongTable(cells,repeatRows=1,colWidths=[(landscape(A4)[0]-48)/len(headers)]*len(headers));table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#d9efea')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,-1),.3,colors.HexColor('#cbd8dc'))]));story.append(table)
    def footer(canvas,doc):canvas.setFont('Helvetica',8);canvas.drawString(24,14,f'FuelFlow Ghana | Page {doc.page}')
    doc.build(story,onFirstPage=footer,onLaterPages=footer);return buffer.getvalue(),'application/pdf'
@router.get('/reports/{kind}/export')
def export(kind:ReportKind,station_id:int,db:DB,actor:Actor,format:Literal['csv','xlsx','pdf']='csv',start:date=Query(default_factory=today),end:date=Query(default_factory=today),nozzle_id:int | None=None,attendant_id:int | None=None,product_id:int | None=None,payment_method:Literal['cash','momo','card','credit'] | None=None):
    permit(actor,'owner','manager','supervisor','accountant','auditor');scoped(actor,station_id)
    data=report_data(db,station_id,kind,start,end,nozzle_id,attendant_id,product_id,payment_method);body,mime=export_bytes(data,format)
    return Response(body,media_type=mime,headers={'Content-Disposition':f'attachment; filename="fuelflow-{kind}-{start}-{end}.{format}"','Cache-Control':'no-store'})
@router.get('/dashboard')
def dashboard(station_id:int,db:DB,actor:Actor):
    permit(actor,'owner','manager','supervisor','accountant','auditor');scoped(actor,station_id)
    summary=day_summary(db,station_id,today());summary['collections']=summary['fuel_collections'];alerts=[]
    for t in db.scalars(select(Tank).where(Tank.station_id==station_id)).all():
        if t.stock<=t.reorder:alerts.append(f'{t.name}: low stock ({t.stock} L)')
    for e in db.scalars(select(Equipment).where(Equipment.station_id==station_id)).all():
        if e.status=='maintenance':alerts.append(f'{e.name}: under maintenance')
        if e.next_service_date and e.next_service_date<=today():alerts.append(f'{e.name}: service due {e.next_service_date}')
    outstanding=ZERO
    for c in db.scalars(select(Customer).where(Customer.station_id==station_id)).all():
        totals=customer_totals(db,c.id);outstanding+=totals['balance']
        if totals['overdue']>0:alerts.append(f'{c.name}: overdue debt GHS {totals["overdue"]}')
    for s in db.scalars(select(Shift).where(Shift.station_id==station_id,Shift.variance!=0,Shift.status.in_(['submitted','reconciled']))).all():alerts.append(f'Shift #{s.id}: variance GHS {s.variance}; review required')
    for i in db.scalars(select(Incident).where(Incident.station_id==station_id,Incident.status=='open')).all():alerts.append(f'Incident #{i.id}: {i.category} ({i.severity})')
    summary.update(alerts=alerts[:100],outstanding_credit=outstanding,active_shifts=db.scalar(select(func.count()).select_from(Shift).where(Shift.station_id==station_id,Shift.status.in_(['open','submitted']))))
    summary['trend']=[day_summary(db,station_id,today()-timedelta(days=i)) for i in range(6,-1,-1)]
    return summary
